# -*- coding: utf-8 -*-
"""
Excel -> ConfigInstance.lua / Language.lua (no CSV).

Sheet layout (row 1-based):
  1: Chinese comments (skipped)
  2: field names (config: col1 = row index, any name; language: Key anywhere; empty name kept as "")
  3: types: number | string | number[] | string[] | table
  4+: data

Rules:
  - number empty -> ERROR (abort)
  - number[] empty cell -> {} ; empty element inside '|' list -> ERROR
  - string[] empty -> {}
  - config mode: row index = first column (name may not be Id); duplicate -> ERROR
  - language mode: index column must be Key; duplicate Key -> ERROR
  - cell '^' -> ',' (legacy escape; real commas in Excel are fine)
  - array elements split by '|'
  - language mode: value cols = language tags like zh-cn / en-us

Requires Python 3.10+ (Path.write_text newline=). Hook skips older interpreters.
"""

from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

LANG_COL_RE = re.compile(r"^[a-z]{2}(-[a-z0-9]+)+$", re.IGNORECASE)
# Path.write_text(..., newline=) is 3.10+; keep in sync with BuildConfig.ps1 / Install-Python.ps1.
MIN_PYTHON = (3, 10)


class ConfigError(Exception):
    """Build abort with file / row / field context."""


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def _apply_caret(text: str) -> str:
    return text.replace("^", ",")


def _lua_string(text: str) -> str:
    text = _apply_caret(str(text))
    text = text.replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n").replace("\r", "")
    return f'"{text}"'


def _to_number(value: Any) -> int | float:
    if _is_empty(value):
        raise ValueError("number is empty")
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    s = _apply_caret(str(value).strip())
    if s == "":
        raise ValueError("number is empty")
    num = float(s)
    if num.is_integer():
        return int(num)
    return num


def _format_number(num: int | float) -> str:
    if isinstance(num, float):
        if math.isnan(num) or math.isinf(num):
            raise ValueError(f"invalid number: {num}")
        if num.is_integer():
            return str(int(num))
        return repr(num)
    return str(num)


def _parse_array(value: Any, elem_type: str) -> str:
    if _is_empty(value):
        return "{}"
    raw = _apply_caret(str(value).strip())
    if raw == "":
        return "{}"
    parts = raw.split("|")
    out: list[str] = []
    for idx, part in enumerate(parts):
        part = part.strip()
        if elem_type == "number":
            if part == "":
                raise ValueError(f"number[] element[{idx}] is empty")
            out.append(_format_number(_to_number(part)))
        else:
            out.append(_lua_string(part))
    return "{" + ",".join(out) + "}"


def _parse_cell(value: Any, type_name: str) -> str:
    t = (type_name or "string").strip().lower()
    if t == "number":
        return _format_number(_to_number(value))
    if t == "string":
        if _is_empty(value):
            return '""'
        return _lua_string(value)
    if t == "number[]":
        return _parse_array(value, "number")
    if t == "string[]":
        return _parse_array(value, "string")
    if t == "table":
        if _is_empty(value):
            return "{}"
        return _apply_caret(str(value).strip())
    raise ValueError(f"Unrecognized type: {type_name}")


def _trim_columns(names: list[Any], types: list[Any]) -> tuple[list[str], list[str]]:
    n = max(len(names), len(types))
    names = list(names) + [None] * (n - len(names))
    types = list(types) + [None] * (n - len(types))
    last = -1
    for i in range(n):
        if names[i] is not None or types[i] is not None:
            last = i
    if last < 0:
        return [], []
    names = names[: last + 1]
    types = types[: last + 1]
    out_names: list[str] = []
    out_types: list[str] = []
    for name, typ in zip(names, types):
        if typ is None and name is None:
            out_names.append("")
            out_types.append("string")
            continue
        if typ is None:
            raise ValueError(f"missing type for field {name!r}")
        out_names.append("" if name is None else str(name).strip())
        out_types.append(str(typ).strip())
    return out_names, out_types


def _find_key_index(names: list[str], *, language_mode: bool) -> int:
    """Language: find Key by name (any column). Config: always first column."""
    if not names:
        raise ValueError("no columns")
    if language_mode:
        for i, name in enumerate(names):
            if (name or "").strip().lower() == "key":
                return i
        raise ValueError(f"no Key column in {names}")
    # Config: index is always column 0, regardless of header text (Id / ItemId / ...).
    return 0


def _normalize_key(value: Any) -> str:
    if _is_empty(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _err(
    file_name: str,
    excel_row: int | None,
    field: str | None,
    detail: str,
    *,
    value: Any = None,
    type_name: str | None = None,
) -> ConfigError:
    lines = [
        "配置校验失败，提交已中止。",
        "",
        f"文件: {file_name}",
    ]
    if excel_row is not None:
        lines.append(f"行号: {excel_row}")
    if field:
        lines.append(f"字段: {field}")
    if type_name:
        lines.append(f"类型: {type_name}")
    if field is not None or excel_row is not None:
        if _is_empty(value):
            lines.append("当前值: (空)")
        else:
            lines.append(f"当前值: {value!r}")
    lines.append(f"原因: {detail}")
    return ConfigError("\n".join(lines))


def _read_sheet(path: Path) -> tuple[list[str], list[str], list[tuple[int, tuple[Any, ...]]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if len(rows) < 3:
        raise _err(path.name, None, None, "表头不足，至少需要 3 行（说明 / 字段名 / 类型）")
    try:
        names, types = _trim_columns(list(rows[1]), list(rows[2]))
    except ValueError as exc:
        raise _err(path.name, 3, None, f"表头类型行无效: {exc}") from exc
    if not names:
        raise _err(path.name, 2, None, "字段名行为空")
    data: list[tuple[int, tuple[Any, ...]]] = []
    width = len(names)
    for excel_row, row in enumerate(rows[3:], start=4):
        if row is None:
            continue
        cells = list(row) + [None] * (width - len(row))
        cells = cells[:width]
        if all(_is_empty(c) for c in cells):
            continue
        data.append((excel_row, tuple(cells)))
    return names, types, data


def _class_name(path: Path) -> str:
    return path.stem.lower() + "Conf"


def _format_key_literal(key_value: Any, key_type: str) -> str:
    t = key_type.strip().lower()
    if t == "number":
        return f"[{_format_number(_to_number(key_value))}]"
    if _is_empty(key_value):
        raise ValueError("行主键为空")
    return f"[{_lua_string(key_value)}]"


def convert_workbook(
    path: Path,
    *,
    language_mode: bool,
) -> tuple[str, list[str], list[str]]:
    """
    Returns (class_name, key_field_names, row_lines like '\\t[id]={...},').
    """
    names, types, data = _read_sheet(path)
    try:
        key_i = _find_key_index(names, language_mode=language_mode)
    except ValueError as exc:
        if language_mode:
            raise _err(path.name, 2, None, f"找不到 Key 列（语言表索引）: {exc}") from exc
        raise _err(path.name, 2, None, f"配置表缺少第一列（索引列）: {exc}") from exc

    key_col_name = (names[key_i] or "").strip() or ("Key" if language_mode else "第一列")
    # Display label in errors: language => Key; config => actual header (or 第一列)
    if language_mode:
        key_label = "Key"
    else:
        key_label = f"第一列({key_col_name})" if key_col_name != "第一列" else "第一列"

    class_name = _class_name(path)

    value_indexes = [i for i in range(len(names)) if i != key_i]
    if language_mode:
        value_indexes = [i for i in value_indexes if LANG_COL_RE.match(names[i] or "")]
        if not value_indexes:
            raise _err(path.name, 2, None, "找不到语言列（如 zh-cn / en-us）")

    key_fields = [names[i] for i in value_indexes]
    seen_keys: dict[str, int] = {}
    row_lines: list[str] = []

    for excel_row, row in data:
        key_raw = row[key_i]
        key_norm = _normalize_key(key_raw)
        if key_norm == "":
            raise _err(
                path.name,
                excel_row,
                key_col_name,
                f"{key_label} 不能为空",
                value=key_raw,
                type_name=types[key_i],
            )

        # Uniqueness: config first-col / language Key
        first_row = seen_keys.get(key_norm)
        if first_row is not None:
            raise _err(
                path.name,
                excel_row,
                key_col_name,
                f"{key_label} 重复: “{key_norm}”（首次出现在第 {first_row} 行）",
                value=key_raw,
                type_name=types[key_i],
            )
        seen_keys[key_norm] = excel_row

        try:
            key_lit = _format_key_literal(key_raw, types[key_i])
        except Exception as exc:  # noqa: BLE001
            raise _err(
                path.name,
                excel_row,
                key_col_name,
                f"{key_label} 解析失败: {exc}",
                value=key_raw,
                type_name=types[key_i],
            ) from exc

        values: list[str] = []
        for i in value_indexes:
            field = names[i] or f"col{i}"
            typ = types[i]
            try:
                values.append(_parse_cell(row[i], typ))
            except Exception as exc:  # noqa: BLE001
                detail = str(exc)
                if "number is empty" in detail:
                    detail = "number 类型不能为空（请填写数字，0 也可以）"
                elif "number[] element" in detail:
                    detail = f"number[] 中存在空元素（用 | 分隔时每一段都要有数字）: {exc}"
                elif "Unrecognized type" in detail:
                    detail = f"无法识别的类型: {typ}"
                raise _err(
                    path.name,
                    excel_row,
                    field,
                    detail,
                    value=row[i],
                    type_name=typ,
                ) from exc
        row_lines.append(f"\t{key_lit}={{{','.join(values)}}},")
    if row_lines:
        row_lines[-1] = row_lines[-1].rstrip(",")
    return class_name, key_fields, row_lines


def build_config_lua(excel_dir: Path) -> str:
    files = sorted(f for f in excel_dir.glob("*.xlsx") if not f.name.startswith("~$"))
    if not files:
        raise ConfigError(f"配置校验失败，提交已中止。\n\n目录: {excel_dir}\n原因: 未找到 xlsx 文件")

    blocks: list[tuple[str, list[str], list[str]]] = []
    for path in files:
        print(f"Conversion profile:{path.name}")
        blocks.append(convert_workbook(path, language_mode=False))

    lines: list[str] = ["", "local ConfigInstance ={}", "ConfigInstance.key={"]
    for class_name, key_fields, _ in blocks:
        field_lit = ",".join(_lua_string(f) for f in key_fields)
        lines.append(f"{class_name}={{{field_lit}}},")
    lines.append("}")
    for class_name, _, row_lines in blocks:
        lines.append("")
        lines.append(f"ConfigInstance.{class_name}={{")
        lines.extend(row_lines)
        lines.append("}")
    lines.append("return ConfigInstance")
    return "\n".join(lines) + "\n"


def build_language_lua(excel_dir: Path) -> str:
    files = sorted(f for f in excel_dir.glob("*.xlsx") if not f.name.startswith("~$"))
    if not files:
        raise ConfigError(f"配置校验失败，提交已中止。\n\n目录: {excel_dir}\n原因: 未找到 xlsx 文件")

    blocks: list[tuple[str, list[str], list[str]]] = []
    for path in files:
        print(f"Conversion profile:{path.name}")
        blocks.append(convert_workbook(path, language_mode=True))

    lines: list[str] = ["", "local Language ={}", "Language.key={"]
    for class_name, key_fields, _ in blocks:
        field_lit = ",".join(_lua_string(f) for f in key_fields)
        lines.append(f"{class_name}={{{field_lit}}},")
    lines.append("}")
    for class_name, _, row_lines in blocks:
        lines.append("")
        lines.append(f"Language.{class_name}={{")
        lines.extend(row_lines)
        lines.append("}")
    lines.append("return Language")
    return "\n".join(lines) + "\n"


def _write_error_file(message: str, kind: str = "") -> None:
    err_dir = Path(__file__).resolve().parent
    err_path = err_dir / "precommit-last-error.txt"
    kind_path = err_dir / "precommit-last-error.kind"
    try:
        # UTF-8 with BOM helps Windows Notepad / some consoles
        err_path.write_bytes(b"\xef\xbb\xbf" + (message.strip() + "\n").encode("utf-8"))
    except OSError:
        pass
    if kind:
        try:
            kind_path.write_bytes((kind.strip() + "\n").encode("ascii"))
        except OSError:
            pass


def _python_too_old_message() -> str:
    found = f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}"
    need = f"{MIN_PYTHON[0]}.{MIN_PYTHON[1]}"
    return (
        "配置校验失败，提交已中止。\n\n"
        f"原因: Python 版本过低（当前 {found}，需要 {need}+）。\n\n"
        "请运行 scripts\\ugit\\2-安装Python环境.bat 安装 Python 3.12 后重试。"
    )


def _write_utf8_lf(path: Path, text: str) -> None:
    """Write UTF-8 text with LF newlines. Path.write_text(newline=) needs 3.10."""
    try:
        path.write_text(text, encoding="utf-8", newline="\n")
        return
    except TypeError:
        # Python < 3.10: open(newline=) still forces LF.
        pass
    with path.open("w", encoding="utf-8", newline="\n") as fh:
        fh.write(text)


def main() -> int:
    if sys.version_info < MIN_PYTHON:
        msg = _python_too_old_message()
        print(f"[ExcelToLua][ERROR]\n{msg}", file=sys.stderr)
        _write_error_file(msg, "PythonTooOld")
        return 1

    parser = argparse.ArgumentParser(description="Excel direct export to Lua config")
    parser.add_argument("--excel-dir", required=True, help="directory of xlsx files")
    parser.add_argument("--out", required=True, help="output .lua path")
    parser.add_argument(
        "--mode",
        choices=("config", "language"),
        default="config",
        help="config => ConfigInstance; language => Language",
    )
    args = parser.parse_args()

    excel_dir = Path(args.excel_dir)
    out_path = Path(args.out)
    if not excel_dir.is_dir():
        msg = f"配置校验失败，提交已中止。\n\n目录: {excel_dir}\n原因: Excel 目录不存在"
        print(f"[ExcelToLua][ERROR]\n{msg}", file=sys.stderr)
        _write_error_file(msg)
        return 1

    print("Start converting files (Excel direct, no CSV)...")
    try:
        if args.mode == "language":
            text = build_language_lua(excel_dir)
        else:
            text = build_config_lua(excel_dir)
    except ConfigError as exc:
        msg = str(exc)
        print(f"[ExcelToLua][ERROR]\n{msg}", file=sys.stderr)
        _write_error_file(msg)
        return 1
    except Exception as exc:  # noqa: BLE001
        msg = f"配置校验失败，提交已中止。\n\n原因: {exc}"
        print(f"[ExcelToLua][ERROR]\n{msg}", file=sys.stderr)
        _write_error_file(msg)
        return 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    _write_utf8_lf(out_path, text)
    err_path = Path(__file__).resolve().parent / "precommit-last-error.txt"
    if err_path.exists():
        try:
            err_path.unlink()
        except OSError:
            pass
    print(f"OK!!!! wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
